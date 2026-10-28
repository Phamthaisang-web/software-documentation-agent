import os
import sys
import docx
import openpyxl
import pytest

AGENT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATES_DIR = os.path.join(AGENT_ROOT, "03_Templates")
TOOLS_DIR = os.path.join(AGENT_ROOT, "09_Tools")
OUTPUT_DIR = os.path.join(AGENT_ROOT, "output")
sys.path.insert(0, TOOLS_DIR)

from document_generator import DocumentGenerator

@pytest.fixture
def sample_profile():
    return {
        "project_id": "test_app",
        "project_name": "Test Application",
        "project_goal": "Hệ thống kiểm thử khả năng bảo tồn định dạng mẫu tài liệu.",
        "functional_requirements": [
            {"id": "REQ-001", "title": "Tính năng đăng nhập", "description": "Xác thực người dùng", "status": "VERIFIED"}
        ],
        "missing_information": []
    }

def test_docx_template_readable_and_structured():
    """Test: All generated DOCX templates can be opened and parsed by python-docx."""
    for root, _, files in os.walk(TEMPLATES_DIR):
        for f in files:
            if f.endswith(".docx"):
                path = os.path.join(root, f)
                doc = docx.Document(path)
                assert len(doc.paragraphs) > 0, f"{path} has no paragraphs"
                assert len(doc.tables) > 0, f"{path} has no tables"

def test_xlsx_template_readable_and_structured():
    """Test: All generated XLSX templates can be opened and parsed by openpyxl."""
    for root, _, files in os.walk(TEMPLATES_DIR):
        for f in files:
            if f.endswith(".xlsx"):
                path = os.path.join(root, f)
                wb = openpyxl.load_workbook(path)
                assert len(wb.sheetnames) >= 1
                ws = wb.active
                assert ws.max_row >= 1
                assert ws.max_column >= 1

def test_generation_does_not_overwrite_source_template(sample_profile):
    """Test: Rendering an output document creates a new file in output/ without modifying template."""
    generator = DocumentGenerator(AGENT_ROOT)
    doc_id = "DOC-PM-001"
    tpl_meta = generator.tpl_registry[generator.doc_registry[doc_id]["preferred_template_id"]]
    tpl_path = os.path.join(AGENT_ROOT, tpl_meta["file_path"])
    
    # Record original modification time
    orig_mtime = os.path.getmtime(tpl_path)
    
    out_file = generator.generate_document(doc_id, sample_profile)
    assert os.path.exists(out_file)
    assert os.path.getsize(out_file) > 0
    assert "output" in out_file

    # Template mtime must not have changed
    new_mtime = os.path.getmtime(tpl_path)
    assert orig_mtime == new_mtime, "Original template was modified during generation!"

def test_generated_output_file_can_be_read_back(sample_profile):
    """Test: Output DOCX and XLSX files can be read back cleanly."""
    generator = DocumentGenerator(AGENT_ROOT)
    
    # DOCX test
    docx_out = generator.generate_document("DOC-REQ-001", sample_profile)
    doc = docx.Document(docx_out)
    full_text = " ".join(p.text for p in doc.paragraphs)
    assert "Test Application" in full_text

    # XLSX test
    xlsx_out = generator.generate_document("DOC-REQ-003", sample_profile)
    wb = openpyxl.load_workbook(xlsx_out)
    assert wb.active.max_row >= 2
